import json, os, shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'SOPs-herzien-Word'
OUT.mkdir(exist_ok=True)
NAVY = '17365D'; BLUE = '2E74B5'; LIGHT = 'E8EEF5'; GRAY = '667085'; WHITE = 'FFFFFF'

EXTRA = {
'SOP-001': dict(trigger='Gebruik deze SOP bij een nieuwe werkwijze, een fout met herhalingsrisico, een wijziging in wet- of klantvereisten, een nieuwe logistieke partner en tijdens de kwartaalreview.', decisions=[('Direct herstellen','Kleine, eenmalige fout met beperkte impact en duidelijke oorzaak.'),('Verbeteractie','Fout kan terugkomen of raakt meerdere orders, producten of partners.'),('Stop/extern advies','Mogelijk productveiligheids-, wettelijk of ernstig financieel risico.')], controls=['Geen actie zonder eigenaar en datum.','Een kritieke actie wordt door de andere compagnon beoordeeld.','Een SOP wordt alleen uitgebreid wanneer de extra stap werkelijk uitvoerbaar is.'], outputs=['Actuele werkwijze','Gesloten actie met bewijs','Kort kwartaalreviewverslag'], example='Een voorraadverschil wordt eerst hersteld. Wanneer dezelfde oorzaak opnieuw voorkomt, wordt niet nogmaals alleen gecorrigeerd: de transacties, fysieke telling en werkwijze van de betrokken locatie worden onderzocht.'),
'SOP-002': dict(trigger='Gebruik vóór eerste inkoop/verkoop van een artikel en bij wijziging van leverancier, verpakking, markt, certificaat, batch-/lotvereiste of houdbaarheid.', decisions=[('Vrijgeven','Verplichte gegevens en documenten zijn compleet en passend.'),('Voorwaardelijk','Niet-kritieke informatie ontbreekt; beperking en einddatum zijn vastgelegd.'),('Blokkeren','Kritisch document, identiteit, productkoppeling of markttoegang is onzeker.')], controls=['SKU, eenheid en documentstatus worden door de andere compagnon gecontroleerd.','Vervaldatum en product-/leverancierskoppeling staan bij elk kritisch document.'], outputs=['Vrijgegeven of geblokkeerd productdossier','Correcte artikelgegevens in Exact Online zodra de inrichting is bevestigd'], example='Een aantrekkelijke partij medische producten wordt niet aangeboden op basis van alleen een prijslijst. Eerst wordt vastgesteld welk exact product, welke fabrikant, welke markt en welke documentatie bij de partij horen.'),
'SOP-003': dict(trigger='Gebruik bij een nieuwe leverancier, een gewijzigde betaalrekening, een kritieke dienstverlener, een grote nieuwe handelsmogelijkheid of terugkerende prestatieproblemen.', decisions=[('Normaal','Vervangbaar, beperkte impact en geen kritieke compliance-afhankelijkheid.'),('Kritisch','Hoge klantimpact, moeilijk vervangbaar of grote product-/financiële afhankelijkheid.'),('Blokkeren','Fraude-indicatie, onbeheerst productrisico of ontbrekende kritieke documentatie.')], controls=['Bankwijzigingen altijd via onafhankelijk contactkanaal bevestigen.','De eerste levering van een nieuwe leverancier extra controleren.'], outputs=['Leveranciersstatus en voorwaarden','Bewijs van verificatie','Acties bij afwijkingen'], example='Een scherpe aanbieding is commercieel interessant, maar leidt nog niet tot een PO. Eerst worden leverancier en product geverifieerd en wordt concrete afzet gezocht.'),
'SOP-004': dict(trigger='Gebruik bij iedere kansgestuurde partij, terugkerende klantvraag, replenishmentbeslissing en voordat een PO wordt bevestigd.', decisions=[('Kansgestuurde handel','Eerst geschikte klanten benaderen en concrete vraag/afzet valideren; daarna inkopen.'),('Herhaalhandel','Vooruit inkopen mag binnen een onderbouwde afzet-, marge- en voorraadlimiet.'),('Niet inkopen','Afzet, marge, documentatie, cashbeslag of logistieke uitvoering is onvoldoende beheerst.')], controls=['Geen aantrekkelijke inkoop verwarren met een verkoopbare deal.','Orderbevestiging vergelijken met PO op prijs, aantal, datum en voorwaarden.','Voorraadlocatie kiezen met het actuele beslismodel intern versus 3PL.'], outputs=['Go/no-go inkoopbesluit','Bevestigde PO','Inboundverwachting en klantafspraak'], example='Bij een aanbieding voor een partij medische producten zoeken Jorn/Hidde eerst afnemers in het bekende netwerk. Alleen de aantoonbaar verkoopbare hoeveelheid plus een bewuste risicobuffer wordt ingekocht.'),
'SOP-005': dict(trigger='Gebruik zodra een leverancier of vervoerder een verwachte aankomst meldt en vóór goederen naar Amstelveen, een extern magazijn / 3PL of rechtstreeks naar een klant vertrekken.', decisions=[('Amstelveen','Klein, kortdurend en aantoonbaar tijdsefficiënt.'),('Extern magazijn / 3PL','Groot volume, opslag, herhaalde handling of schaal-/foutrisico.'),('Direct naar klant','Tussenopslag voegt geen waarde toe en documenten/traceerbaarheid zijn geborgd.')], controls=['Leverdatum en tijdvak schriftelijk bevestigen.','Bij uitblijven leverancier vóór afgesproken tijd direct actief nabellen en alternatief plannen.'], outputs=['Volledige vooraankondiging','Bevestigde bestemming en ETA','Afwijking/escalatie bij vertraging'], example='Een leverancier die “in de ochtend” belooft te leveren krijgt een concreet tijdvak en contactpersoon. Bij no-show volgt niet pas om 16:00 uur actie; de eerste escalatie staat vooraf gepland.'),
'SOP-006': dict(trigger='Gebruik bij iedere fysieke ontvangst, inclusief leverancier, retour en transfer.', decisions=[('Vrijgeven','Artikel, aantal, staat en vereiste gegevens zijn akkoord.'),('Quarantaine','Schade, mismatch, documentgebrek of twijfel.'),('Weigeren','Onveilig, niet identificeerbaar of niet herstelbaar zonder onbeheerst risico.')], controls=['3PL levert bij verschil foto’s, telling, locatie en transactiebewijs.','Voorraad wordt pas beschikbaar na fysieke bevestiging en statusbesluit.'], outputs=['Ontvangstbewijs','Juiste voorraadboeking','Afwijkingsdossier indien nodig'], example='Bij één ontbrekend collo controleert de 3PL niet alleen de loslijst, maar ook fysieke locatie, scans en eventueel camerabeeld of overdrachtsbewijs.'),
'SOP-007': dict(trigger='Gebruik bij iedere nieuwe inbound, opslagvraag, transfer, retour, spoedorder of order waarbij intern uitvoeren wordt overwogen.', decisions=[('Intern','Klein, snel, geen structurele opslag en het aantoonbare voordeel overstijgt de opportunity cost.'),('Extern magazijn / 3PL','Palletvolume, opslag, veel handelingen, foutkans of verlies van commerciële tijd.'),('3PL-keuze herzien','Huidige partner presteert structureel onvoldoende of kandidaat biedt aantoonbaar betere totale waarde.')], controls=['Gebruik altijd het actuele Excel-beslismodel bij twijfel.','Logicall Zaandam blijft kandidaat totdat offerte, portal-demo, SLA, testflow en migratieplan zijn beoordeeld.','Scan Global Logistics wordt als huidige 3PL beoordeeld op feiten, niet automatisch behouden of afgewezen.'], outputs=['Vastgelegde locatiekeuze','Kosten- en tijdsvergelijking','Transfer- of migratieplan waar nodig'], example='Zelf ontvangen en verplaatsen lijkt goedkoop, maar 3 uur van een compagnon plus ritten en foutkans kan duurder zijn dan volledige 3PL-handling. “Het kan intern” is daarom geen besluitcriterium.'),
'SOP-008': dict(trigger='Gebruik bij een geplande telling, voorraadverschil, vermiste voorraad, schade, verlopen product of incourante voorraad.', decisions=[('Direct corrigeren','Klein, verklaarbaar en volledig bewezen binnen afgesproken grens.'),('Hertellen/onderzoek','Materieel, onverklaarbaar of mogelijk transactiefout.'),('Blokkeren/afwaarderen','Niet verkoopbaar, verlopen, beschadigd of onzeker.')], controls=['De uitvoerende 3PL is actief verantwoordelijk voor bewijs en hertelling.','Materiële correctie wordt door de andere compagnon goedgekeurd.'], outputs=['Tel- en correctiebewijs','Juiste voorraad en waarde','Verbeteractie of claim'], example='Een vermiste pallet wordt niet administratief weggeboekt omdat het portal onduidelijk is. Eerst volgen WMS-historie, locatiecontrole, hertelling en overdrachtsbewijs.'),
'SOP-009': dict(trigger='Gebruik voor iedere klantorder, inclusief kansgestuurde verkoop, herhaalorder, spoedorder en sample op klantverzoek.', decisions=[('Vrijgeven','Klant, product, prijs, voorraad, voorwaarden en compliance zijn duidelijk.'),('Hold','Kritieke informatie of voorraadzekerheid ontbreekt.'),('Afwijzen/heronderhandelen','Marge, risico, betaalconditie of uitvoering is niet verantwoord.')], controls=['Samples alleen op klantverzoek en waar mogelijk via de leverancier laten meelopen.','Afwijkende prijs of risico door de andere compagnon laten controleren.'], outputs=['Bevestigde klantorder','Duidelijke uitvoeringsinstructie','Backorder- of holdstatus'], example='Bij een kansgestuurde partij worden geïnteresseerde klanten eerst concreet bevraagd. Pas na bevestigde vraag wordt de inkoophoeveelheid definitief gemaakt.'),
'SOP-010': dict(trigger='Gebruik zodra een klantorder is vrijgegeven voor fysieke uitvoering.', decisions=[('Standaardcontrole','Reguliere, bekende order en geautomatiseerde/3PL-uitvoering.'),('Tweede controle','Handmatige, kritieke, hoge waarde of afwijkende zending.'),('Stop','Mismatch, schade, onduidelijk adres of ontbrekend document.')], controls=['Geen alternatief artikel of batch zonder akkoord.','Dispatchbewijs bevat tracking en aantal colli.'], outputs=['Correct verpakte zending','Dispatchbevestiging','Afwijkingsmelding indien nodig'], example='Een spoedorder wordt wel geprioriteerd, maar niet zonder adres-, artikel-, aantal- en documentcontrole overgedragen.'),
'SOP-011': dict(trigger='Gebruik bij iedere transportboeking en zodra een geplande levering of ophaling afwijkt.', decisions=[('Exception-based volgen','Standaardzending met betrouwbare tracking.'),('Actief volgen','Spoed, hoge waarde, export of eerdere afwijking.'),('Alternatief inzetten','Leverancier/vervoerder mist afgesproken tijd en klantimpact wordt reëel.')], controls=['Werk met concreet tijdvak, cut-off en escalatiemoment.','Bij no-show niet passief wachten tot einde dag.','POD, scans en foto’s direct veiligstellen bij vermissing of schade.'], outputs=['Bevestigde booking en tracking','Proactieve klantupdate','Claimdossier bij schade/verlies'], example='Als een leverancier ’s ochtends moet leveren, wordt vooraf bepaald wanneer wordt nagebeld en wanneer alternatief transport of klantcommunicatie start.'),
'SOP-012': dict(trigger='Gebruik voor iedere internationale zending met commerciële, douane-, oorsprongs- of productspecifieke documentatie.', decisions=[('Vrijgeven','Alle gegevens zijn consistent en markttoegang is duidelijk.'),('Hold','Kritisch document of gegeven ontbreekt.'),('Extern bevestigen','Onbekende douane- of regelgevingseis.')], controls=['Eén definitieve documentset gebruiken.','Nieuwe markt of gereguleerd product door de andere compagnon laten controleren.'], outputs=['Consistente exportdocumentset','Vrijgave of hold','Archiefbewijs'], example='Een HS-code wordt niet blind uit een oude factuur gekopieerd wanneer product of verpakking is gewijzigd.'),
'SOP-013': dict(trigger='Gebruik bij klantklacht, leveranciersfout, transportschade, vermissing, 3PL-afwijking of terugkerende servicefout.', decisions=[('Direct oplossen','Impact beheersen en oorzaak duidelijk.'),('Claim','Derde partij is aantoonbaar verantwoordelijk en bewijs is compleet.'),('Structurele verbetering','Kritieke of herhaalde oorzaak.')], controls=['Containment vóór schuldvraag.','Partner levert objectief bewijs en eigen correctieactie.','Niet sluiten voordat voorraad, klant en financiën zijn verwerkt.'], outputs=['Oplossing en communicatie','Claim/credit','Verbeteractie bij herhaling'], example='Bij transportschade worden verpakking, foto’s bij ontvangst, vervoerdersscan en klantimpact verzameld voordat verantwoordelijkheid wordt vastgesteld.'),
'SOP-014': dict(trigger='Gebruik bij kwartaalreview van de 3PL, materiële incidenten, contractherziening en beoordeling van een mogelijke nieuwe partner.', decisions=[('Behouden','Prestatie en totale waarde zijn voldoende.'),('Voorwaardelijk','Verbeterplan met meetbare deadline nodig.'),('Migratie onderzoeken','Structurele tekortkomingen en aantoonbaar betere kandidaat.'),('Overstappen','Offerte, SLA, IT/portal, testflow, inventarisatie en terugvalplan zijn akkoord.')], controls=['Scan Global Logistics als huidige partner feitelijk meten.','Logicall Zaandam alleen beoordelen op bevestigde offerte, demo en afspraken.','Migratiekosten en uren van Jorn/Hidde meenemen.'], outputs=['Partnerstatus','Actieplan','Onderbouwd migratiebesluit indien relevant'], example='Een mooier portal is waardevol, maar geen zelfstandig migratiebesluit. Ook tarieven, voorraadreconciliatie, medische producthandling, cut-offs en foutafhandeling moeten aantoonbaar beter of beheersbaar zijn.'),
'SOP-015': dict(trigger='Gebruik vanaf ieder retourverzoek, geweigerde levering of teruggekomen sample.', decisions=[('Opnieuw verkoopbaar','Identiteit, staat, opslag en traceerbaarheid zijn aantoonbaar intact.'),('Geblokkeerd/herstel','Inspectie, herverpakking of informatie nodig.'),('Afschrijven/vernietigen','Niet veilig, verlopen, beschadigd of economisch onverantwoord.')], controls=['Nooit direct terugboeken in beschikbare voorraad.','Vrijgave en materiële afschrijving door de andere compagnon controleren.'], outputs=['RMA- en ontvangstbewijs','Dispositiebesluit','Credit/claim en juiste voorraad'], example='Een ongeopende doos is niet automatisch verkoopbaar wanneer opslagcondities of batchtraceerbaarheid tijdens de retourperiode onbekend zijn.')
}

def font(run, size=11, bold=False, color='20242A', italic=False):
    run.font.name='Calibri'; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Calibri'); run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Calibri')
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=RGBColor.from_string(color)

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd')) or OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd) if shd.getparent() is None else None

def set_cell_width(cell, dxa):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW')) or OxmlElement('w:tcW'); tcW.set(qn('w:w'),str(dxa)); tcW.set(qn('w:type'),'dxa'); tcPr.append(tcW) if tcW.getparent() is None else None

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+tag)) or OxmlElement('w:'+tag); node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa'); mar.append(node) if node.getparent() is None else None

def set_table_geometry(table, widths):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr
    tblW=tblPr.find(qn('w:tblW')); tblW.set(qn('w:w'),str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    tblInd=tblPr.find(qn('w:tblInd')) or OxmlElement('w:tblInd'); tblInd.set(qn('w:w'),'120'); tblInd.set(qn('w:type'),'dxa'); tblPr.append(tblInd) if tblInd.getparent() is None else None
    grid=table._tbl.tblGrid
    for x in list(grid): grid.remove(x)
    for width in widths:
        col=OxmlElement('w:gridCol'); col.set(qn('w:w'),str(width)); grid.append(col)
    for row in table.rows:
        for i,cell in enumerate(row.cells): set_cell_width(cell,widths[i]); set_cell_margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_numbering(doc):
    part=doc.part.numbering_part; root=part.element
    abs_id=max([int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum'))] or [0])+1
    num_id=max([int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num'))] or [0])+1
    abstract=OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'),str(abs_id))
    multi=OxmlElement('w:multiLevelType'); multi.set(qn('w:val'),'singleLevel'); abstract.append(multi)
    lvl=OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'),'0'); abstract.append(lvl)
    start=OxmlElement('w:start'); start.set(qn('w:val'),'1'); lvl.append(start)
    numFmt=OxmlElement('w:numFmt'); numFmt.set(qn('w:val'),'decimal'); lvl.append(numFmt)
    lvlText=OxmlElement('w:lvlText'); lvlText.set(qn('w:val'),'%1.'); lvl.append(lvlText)
    suff=OxmlElement('w:suff'); suff.set(qn('w:val'),'tab'); lvl.append(suff)
    pPr=OxmlElement('w:pPr'); tabs=OxmlElement('w:tabs'); tab=OxmlElement('w:tab'); tab.set(qn('w:val'),'num'); tab.set(qn('w:pos'),'540'); tabs.append(tab); pPr.append(tabs)
    ind=OxmlElement('w:ind'); ind.set(qn('w:left'),'540'); ind.set(qn('w:hanging'),'270'); pPr.append(ind); lvl.append(pPr)
    root.append(abstract); num=OxmlElement('w:num'); num.set(qn('w:numId'),str(num_id)); aid=OxmlElement('w:abstractNumId'); aid.set(qn('w:val'),str(abs_id)); num.append(aid); root.append(num)
    return num_id

def numbered(doc,text,num_id):
    p=doc.add_paragraph(style='SOP Step'); pPr=p._p.get_or_add_pPr(); numPr=OxmlElement('w:numPr'); ilvl=OxmlElement('w:ilvl'); ilvl.set(qn('w:val'),'0'); nid=OxmlElement('w:numId'); nid.set(qn('w:val'),str(num_id)); numPr.extend([ilvl,nid]); pPr.append(numPr); font(p.add_run(text)); return p

def bullet(doc,text):
    p=doc.add_paragraph(style='SOP Bullet'); pPr=p._p.get_or_add_pPr(); ind=pPr.find(qn('w:ind')) or OxmlElement('w:ind'); ind.set(qn('w:left'),'540'); ind.set(qn('w:hanging'),'270'); pPr.append(ind) if ind.getparent() is None else None
    r=p.add_run('•\t'+text); font(r); return p

def heading(doc,text):
    p=doc.add_paragraph(style='Heading 1'); p.paragraph_format.keep_with_next=True; font(p.add_run(text),16,True,BLUE); return p

def build(s):
    doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1); sec.header_distance=sec.footer_distance=Inches(.492)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
    for name,size,before,after,color in [('Heading 1',16,18,10,BLUE),('Heading 2',13,14,7,BLUE),('Heading 3',12,10,5,NAVY)]:
        st=doc.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    for name in ['SOP Step','SOP Bullet']:
        st=doc.styles.add_style(name,1); st.font.name='Calibri'; st.font.size=Pt(11); st.paragraph_format.space_after=Pt(4); st.paragraph_format.line_spacing=1.25
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.LEFT; font(hp.add_run('INCO-SOURCE  |  STANDARD OPERATING PROCEDURE'),9,True,GRAY)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run(f"{s['sopNumber']}  |  Versie {s['version']}  |  Beheerst document"),8,False,GRAY)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); font(p.add_run(s['sopNumber']),10,True,BLUE)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(12); p.paragraph_format.keep_with_next=True; font(p.add_run(s['title']),24,True,NAVY)
    t=doc.add_table(rows=5,cols=2); t.style='Table Grid'; set_table_geometry(t,[2700,6660])
    meta=[('Status',s['status']),('Versie',s['version']),('Eigenaar',s['owner']),('Categorie',s['category']),('Beoordeling',s['reviewDate'])]
    for row,(label,val) in zip(t.rows,meta):
        set_cell_shading(row.cells[0],LIGHT); font(row.cells[0].paragraphs[0].add_run(label),9,True,NAVY); font(row.cells[1].paragraphs[0].add_run(val),9)
    extra=EXTRA[s['sopNumber']]
    heading(doc,'1. Doel'); doc.add_paragraph(s['purpose'])
    heading(doc,'2. Operationele context en gebruiksmoment'); doc.add_paragraph(extra['trigger'])
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); font(p.add_run('Waarom dit ertoe doet: '),11,True,NAVY); font(p.add_run(extra['example']))
    heading(doc,'3. Scope'); doc.add_paragraph(s['scope'])
    heading(doc,'4. Rollen en verantwoordelijkheden'); [bullet(doc,x) for x in s['roles']]
    heading(doc,'5. Benodigde systemen en documenten')
    p=doc.add_paragraph(); font(p.add_run('Systemen: '),11,True,NAVY); font(p.add_run('; '.join(s['systems'])))
    p=doc.add_paragraph(); font(p.add_run('Documenten en bewijs: '),11,True,NAVY); font(p.add_run('; '.join(s['documents'])))
    heading(doc,'6. Belangrijkste beslisregels')
    dt=doc.add_table(rows=1,cols=2); dt.style='Table Grid'; set_table_geometry(dt,[2700,6660]); hdr=dt.rows[0].cells
    for c,txt in zip(hdr,['Besluit','Wanneer toepassen']): set_cell_shading(c,LIGHT); font(c.paragraphs[0].add_run(txt),9,True,NAVY)
    for label,detail in extra['decisions']:
        cells=dt.add_row().cells; set_table_geometry(dt,[2700,6660]); font(cells[0].paragraphs[0].add_run(label),9,True,NAVY); font(cells[1].paragraphs[0].add_run(detail),9)
    heading(doc,'7. Werkwijze'); num_id=add_numbering(doc); [numbered(doc,x,num_id) for x in s['steps']]
    heading(doc,'8. Kritieke controles'); [bullet(doc,x) for x in extra['controls']]
    heading(doc,'9. Uitzonderingen en beslismomenten'); [bullet(doc,x) for x in s['exceptions']]
    heading(doc,'10. Escalaties'); [bullet(doc,x) for x in s['escalations']]
    heading(doc,'11. Vereiste output en bewijs'); [bullet(doc,x) for x in extra['outputs']]
    heading(doc,'12. Prestatie-indicatoren'); [bullet(doc,x) for x in s['kpis']]
    heading(doc,'13. Veelgemaakte fouten'); [bullet(doc,x) for x in s['commonMistakes']]
    heading(doc,'14. Gerelateerde SOP’s'); [bullet(doc,x) for x in s['relatedSops']]
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(16); font(p.add_run('Documentbeheersing: alleen de actuele, goedgekeurde versie op de centrale locatie is geldig.'),9,False,GRAY,True)
    fn=f"{s['sopNumber']} {s['title'].replace('/','-')}.docx"; doc.save(OUT/fn)

if __name__=='__main__':
    data=json.loads((ROOT/'data'/'sops.json').read_text(encoding='utf-8'))
    for old in OUT.glob('*.docx'): old.unlink()
    for s in data: build(s)
    print(f'Built {len(data)} DOCX files in {OUT}')
