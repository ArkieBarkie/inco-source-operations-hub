import json
import re
import unicodedata
from pathlib import Path
from docx import Document

ROOT = Path(r"C:\Users\erik-\OneDrive\Documenten\Inco-Source")
SOURCE = ROOT / "Final SOP Inco-Source"
OUTPUT = ROOT / "data" / "sops.json"

PROCESS_MAP = {
    "Besturing & kwaliteit": "Kwaliteitsbeheer",
    "Product & compliance": "Product- en compliancebeheer",
    "Inkoop & leveranciers": "Inkoop en leveranciers",
    "Goederen & voorraad": "Goederen en voorraad",
    "Logistiek & fulfilment": "Order en logistiek",
}

def clean(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()

def slugify(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", value).strip("-")

def rows(table):
    return [[clean(cell.text) for cell in row.cells] for row in table.rows]

def metadata(table):
    result = {}
    for row in rows(table):
        for i in range(0, len(row) - 1, 2):
            result[row[i]] = row[i + 1]
    return result

def after_label(text: str, label: str) -> str:
    return clean(text.split(label, 1)[1]) if label in text else clean(text)

def split_list(text: str):
    return [clean(x) for x in re.split(r";", text) if clean(x)]

def table_pairs(table, separator=" - "):
    result = []
    for row in rows(table)[1:]:
        if len(row) >= 2 and any(row):
            result.append(f"{row[0]}{separator}{row[1]}")
    return result

def import_doc(path: Path):
    doc = Document(path)
    paras = [clean(p.text) for p in doc.paragraphs if clean(p.text)]
    sop_number = re.search(r"SOP-\d{3}", " ".join(paras)).group(0)
    title = paras[1] if len(paras) > 1 else path.stem
    meta = metadata(doc.tables[0])
    status = meta.get("Status", "Vervallen" if "VERVALLEN" in path.name.upper() else "Actief")
    category = meta.get("Categorie", "Overig")
    revision = rows(doc.tables[-1])[-1] if len(doc.tables) > 1 else []
    version = revision[0] if revision and re.match(r"\d", revision[0]) else meta.get("Versie", "1.0")
    last_updated = (revision[1] if len(revision) > 1 else meta.get("Ingangsdatum", "")).split(" ")[0]
    if re.match(r"\d{2}-\d{2}-\d{4}", last_updated):
        d, m, y = last_updated.split("-")
        last_updated = f"{y}-{m}-{d}"
    review = meta.get("Beoordeling", "Niet van toepassing")
    if re.match(r"\d{2}-\d{2}-\d{4}", review):
        d, m, y = review.split("-")
        review = f"{y}-{m}-{d}"

    if status.lower() == "vervallen" or len(doc.tables) < 7:
        reason = revision[2] if len(revision) > 2 else "Deze SOP is vervallen."
        return {
            "id": sop_number.lower(), "sopNumber": sop_number, "title": title,
            "slug": slugify(f"{sop_number} {title}"), "category": category,
            "process": PROCESS_MAP.get(category, category), "summary": reason,
            "purpose": reason, "scope": "Niet van toepassing; raadpleeg de vervangende SOP.",
            "owner": meta.get("Eigenaar", "Jorn / Hidde"), "department": "Inco-Source",
            "roles": [], "systems": [], "documents": [], "steps": [], "exceptions": [],
            "escalations": [], "kpis": [], "commonMistakes": [],
            "relatedSops": re.findall(r"SOP-\d{3}", reason), "status": "Vervallen",
            "version": version, "lastUpdated": last_updated, "reviewDate": review,
            "keywords": [sop_number.lower(), "vervallen", "vervangen"],
            "sourceFile": path.name, "fullText": "\n".join(paras + [" | ".join(r) for t in doc.tables for r in rows(t)]),
            "sourceSections": [{"title": "Vervallenverklaring", "content": [reason]}],
        }

    purpose = after_label(paras[5], "Doel:")
    start_end = paras[6]
    scope = after_label(paras[7], "Scope:")
    headline = after_label(paras[3], "Hoofdregel")
    requirements = table_pairs(doc.tables[1], ": ")
    definitions = table_pairs(doc.tables[2], ": ")
    step_rows = rows(doc.tables[3])[1:]
    steps = [f"{r[2]} - {r[3]} Output: {r[4]}" for r in step_rows if len(r) >= 5]
    roles = []
    for r in step_rows:
        if len(r) >= 4:
            item = f"{r[1]} - verantwoordelijk voor: {r[2].lower()}"
            if item not in roles:
                roles.append(item)
    decisions = table_pairs(doc.tables[4], ": ")
    checks = [r[1] for r in rows(doc.tables[5])[1:] if len(r) > 1 and r[1]]
    exceptions = table_pairs(doc.tables[6], ": ")
    dossier_para = next((p for p in paras if p.startswith("Verplicht dossier:")), "")
    kpi_para = next((p for p in paras if p.startswith("KPI-sturing:")), "")
    related_para = next((p for p in paras if p.startswith("Gerelateerde SOP")), "")
    documents = split_list(after_label(dossier_para, "Verplicht dossier:"))
    kpis = split_list(after_label(kpi_para, "KPI-sturing:"))
    related = [x for x in re.findall(r"SOP-\d{3}", related_para) if x != sop_number]
    source_sections = [
        {"title": "Doel", "content": [purpose]},
        {"title": "Startpunt en eindpunt", "content": [start_end]},
        {"title": "Scope", "content": [scope]},
        {"title": "Vereisten", "content": requirements},
        {"title": "Definities en afkortingen", "content": definitions},
        {"title": "Processtappen", "content": steps},
        {"title": "Belangrijkste beslisregels", "content": decisions},
        {"title": "Verplichte controle", "content": checks},
        {"title": "Uitzonderingen en directe acties", "content": exceptions},
        {"title": "Dossier", "content": documents},
        {"title": "KPI-sturing", "content": kpis},
    ]
    full_parts = paras + [" | ".join(r) for table in doc.tables for r in rows(table)]
    keywords = sorted(set(re.findall(r"[A-Za-zÀ-ÿ0-9/-]{4,}", " ".join(full_parts).lower())))
    return {
        "id": sop_number.lower(), "sopNumber": sop_number, "title": title,
        "slug": slugify(f"{sop_number} {title}"), "category": category,
        "process": PROCESS_MAP.get(category, category), "summary": headline,
        "purpose": purpose, "scope": scope, "owner": meta.get("Eigenaar", "Jorn / Hidde"),
        "department": "Inco-Source", "roles": roles, "systems": [], "documents": documents,
        "steps": steps, "exceptions": exceptions, "escalations": decisions, "kpis": kpis,
        "commonMistakes": checks, "relatedSops": related, "status": status,
        "version": version, "lastUpdated": last_updated, "reviewDate": review,
        "keywords": keywords, "sourceFile": path.name, "fullText": "\n".join(full_parts),
        "sourceSections": source_sections,
    }

items = [import_doc(path) for path in sorted(SOURCE.glob("SOP-*.docx"))]
if len(items) != 15:
    raise SystemExit(f"Expected 15 SOP files, found {len(items)}")
numbers = [item["sopNumber"] for item in items]
if len(numbers) != len(set(numbers)):
    raise SystemExit("Duplicate SOP numbers found")
OUTPUT.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Imported {len(items)} SOPs: {sum(x['status']=='Actief' for x in items)} active, {sum(x['status']=='Vervallen' for x in items)} expired")
